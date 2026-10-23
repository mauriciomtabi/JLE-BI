"""Consolida faturamento ERP e saidas bancarias JLE, sem alterar a planilha."""
import argparse
import datetime as dt
import json
import re
import shutil
import unicodedata
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
CACHE = ROOT / 'executivo_jle_local.xlsx'
OUTPUT = ROOT / 'executivo_jle_data.js'
FOLDERS = [
    Path(r'\\10.121.21.252\financeiro\Angelita\2026\FLUXO DIARIO'),
    Path(r'\\10.121.21.252\financeiro\Angelita\2026'),
    Path(r'\\10.121.21.252\financeiro\Angelita\2026\FLUXO DIARIO\PLANILHAS ANTIGAS'),
]
MONTHS = ['janeiro','fevereiro','marco','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(value or '')).lower() if not unicodedata.combining(c))


def key(value):
    return re.sub('[^a-z0-9]', '', normalize(value))


def cents(value):
    if value is None or value == '':
        return 0
    text = str(value).strip().replace('R$', '').replace(' ', '')
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    number = Decimal(text)
    if not number.is_finite():
        raise ValueError('Valor financeiro nao finito.')
    return int((number * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def date(value):
    if isinstance(value, dt.datetime): return value.date()
    if isinstance(value, dt.date): return value
    if isinstance(value, (int, float)):
        return openpyxl.utils.datetime.from_excel(value).date()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d/%m/%Y %H:%M:%S'):
        try: return dt.datetime.strptime(str(value).strip(), fmt).date()
        except ValueError: pass
    raise ValueError(f'Data nao reconhecida: {value!r}')


def choose_source():
    candidates = []
    for folder in FOLDERS:
        if not folder.is_dir(): continue
        candidates.extend(p for p in folder.glob('*.xlsx') if not p.name.startswith('~$')
                          and 'fluxo de caixa anal' in normalize(p.name)
                          and 'atualizado' in normalize(p.name) and 'drill' not in normalize(p.name))
    def priority(p):
        m=re.search(r'(\d{2})[-.](\d{2})[-.](\d{4})',p.name)
        reference=dt.datetime(int(m[3]),int(m[2]),int(m[1])) if m else dt.datetime.min
        return reference,p.stat().st_mtime
    if candidates: return max(candidates,key=priority),False
    if CACHE.is_file(): return CACHE,True
    raise FileNotFoundError('Planilha JLE indisponivel na rede e no cache executivo.')


def consolidate(source, original_name=None, fallback=False):
    wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
    try:
        receita=next((s for s in wb if key(s.title)=='receita'),None)
        if receita is None: raise ValueError('Aba RECEITA ausente; base anterior preservada.')
        monthly=defaultdict(lambda:{'gross':0,'withholding':0,'net':0,'expense':0,'revenue_count':0,
                                   'expense_count':0,'application':0,'application_count':0,
                                   'transfer_excluded':0,'transfer_count':0,
                                   'initial_excluded':0,'by_uf':defaultdict(int),'by_customer':defaultdict(int),
                                   'by_bank':defaultdict(int),'expense_source_month':False})
        revenue_headers={key(v):i for i,v in enumerate(next(receita.iter_rows(values_only=True)))}
        def rev(row,name): return row[revenue_headers[key(name)]]
        for required in ('Dt.fatur','Valor Bruto','Retenção','Valor Líquido','Razão social','Posição detalhada','UF','Série','Nota','Cli/for'):
            if key(required) not in revenue_headers: raise ValueError(f'Coluna RECEITA ausente: {required}')
        ids=set(); canceled=0; last_invoice=None; records=[]
        for rowno,row in enumerate(receita.iter_rows(min_row=2,values_only=True),2):
            if not rev(row,'Dt.fatur'): continue  # Linhas de resumo nao sao notas.
            status=normalize(rev(row,'Posição detalhada'))
            if 'cancel' in status: canceled+=1; continue
            if 'faturad' not in status: raise ValueError(f'Situacao ERP desconhecida na linha {rowno}.')
            day=date(rev(row,'Dt.fatur')); period=day.strftime('%Y-%m')
            identity=(str(rev(row,'Série')),str(rev(row,'Nota')),str(rev(row,'Cli/for')),str(rev(row,'UF')),str(day))
            if identity in ids: raise ValueError(f'Nota ERP duplicada na linha {rowno}; conferir a fonte.')
            ids.add(identity)
            gross,withholding,net=(cents(rev(row,k)) for k in ('Valor Bruto','Retenção','Valor Líquido'))
            if abs(gross-withholding-net)>1: raise ValueError(f'Bruto - retencoes nao concilia na linha {rowno}.')
            m=monthly[period];m['gross']+=gross;m['withholding']+=withholding;m['net']+=net;m['revenue_count']+=1
            m['by_uf'][str(rev(row,'UF') or 'N/D').split('-')[0].strip()]+=gross
            m['by_customer'][str(rev(row,'Razão social') or 'N/D')]+=gross
            records.append({'kind':'revenue','period':period,'date':day.isoformat(),
                            'customer':str(rev(row,'Razão social') or 'N/D'),
                            'uf':str(rev(row,'UF') or 'N/D').split('-')[0].strip(),
                            'description':'Nota '+str(rev(row,'Nota')),'category':'Faturamento ERP',
                            'bank':'','value':gross/100,'net':net/100,'withholding':withholding/100})
            last_invoice=max(last_invoice,day) if last_invoice else day
        if not ids: raise ValueError('Nenhuma nota faturada valida encontrada.')
        years={int(k[:4]) for k in monthly}; bank_sheets=[]; mismatch_count=0; last_expense=None
        for sheet in wb:
            n=normalize(sheet.title)
            if not any(s in n for s in ('bradesco','confian','maxcredito','maxicredito')): continue
            y=re.search(r'20\d{2}',sheet.title)
            if not y or int(y[0]) not in years: continue
            bank='Bradesco' if 'bradesco' in n else 'Sicoob Confiança' if 'confian' in n else 'Sicoob MaxiCrédito'
            headers=None; found=False
            for rowno,row in enumerate(sheet.iter_rows(values_only=True),1):
                if headers is None:
                    mapping={key(v):i for i,v in enumerate(row) if v is not None}
                    if 'data' in mapping and 'entradasaida' in mapping:
                        headers=mapping;found=True
                    elif rowno>=15: break
                    continue
                def get(name):
                    i=headers.get(key(name));return row[i] if i is not None and i<len(row) else None
                if not get('Data'): continue
                flow=key(get('Entrada / Saída'))
                if flow not in ('entrada','saida'): continue
                day=date(get('Data'));period=day.strftime('%Y-%m')
                competence=normalize(get('Competência')).strip()
                if competence:
                    recognized=next((i+1 for i,v in enumerate(MONTHS) if competence.split('/')[0].strip()==v),None)
                    if recognized:
                        cy=re.search(r'20\d{2}',competence)
                        period=f'{int(cy[0]) if cy else day.year:04d}-{recognized:02d}'
                if period!=day.strftime('%Y-%m'): mismatch_count+=1
                m=monthly[period];m['expense_source_month']=True
                if flow!='saida': continue
                category=get('Movimento') or get('Receitas / Despesas') or get('Fornecedores') or get('D') or ''
                description=get('Descrição') or get('Detalhamento') or get('Coluna1') or ''
                value=abs(cents(get('R$ (Valores)')))
                if not value: continue
                if 'saldoinicial' in key(category): m['initial_excluded']+=value;continue
                # Meio de pagamento "transferencia bancaria" sozinho NAO e exclusao.
                classification=key(str(category)+' '+str(description))
                if 'transferenciaentreconta' in classification or key(category) in ('transferencia','transferencias','transferenciainterna'):
                    m['transfer_excluded']+=value;m['transfer_count']+=1;continue
                
                # Aplicações financeiras na coluna de movimentos não entram como despesa
                mov_val = str(get('Movimento') or category or '')
                if 'aplicac' in key(mov_val) or key(mov_val).startswith('aplic'):
                    m['application']+=value;m['application_count']+=1
                    records.append({'kind':'application','period':period,'date':day.isoformat(),'customer':'',
                                    'uf':str(get('UF') or '').strip().upper() if str(get('UF') or '').strip().upper() in {'AC','AL','AP','AM','BA','CE','DF','ES','GO','MA','MT','MS','MG','PA','PB','PR','PE','PI','RJ','RN','RS','RO','RR','SC','SP','SE','TO'} else 'N/D',
                                    'description':str(description),'category':str(category) if category else 'Aplicação','bank':bank,'value':value/100})
                    last_expense=max(last_expense,day) if last_expense else day
                    continue

                m['expense']+=value;m['expense_count']+=1;m['by_bank'][bank]+=value
                records.append({'kind':'expense','period':period,'date':day.isoformat(),'customer':'',
                                'uf':str(get('UF') or '').strip().upper() if str(get('UF') or '').strip().upper() in {'AC','AL','AP','AM','BA','CE','DF','ES','GO','MA','MT','MS','MG','PA','PB','PR','PE','PI','RJ','RN','RS','RO','RR','SC','SP','SE','TO'} else 'N/D',
                                'description':str(description),'category':str(category),'bank':bank,'value':value/100})
                last_expense=max(last_expense,day) if last_expense else day
            if found: bank_sheets.append(sheet.title.strip())
        if not bank_sheets: raise ValueError('Nenhuma aba bancaria valida encontrada.')
        def money(v):return v/100
        result=[]
        for period,m in sorted(monthly.items()):
            comparable=m['revenue_count']>0 and m['expense_source_month']
            result.append({'period':period,'revenue':money(m['gross']) if m['revenue_count'] else None,
                           'expense':money(m['expense']) if m['expense_source_month'] else None,
                           'application':money(m['application']),
                           'application_count':m['application_count'],
                           'result':money(m['gross']-m['expense']) if comparable else None,
                           'comparable':comparable,'revenue_count':m['revenue_count'],'expense_count':m['expense_count'],
                           'withholding':money(m['withholding']),'net_revenue':money(m['net']),
                           'transfers_excluded':money(m['transfer_excluded']),'transfer_count':m['transfer_count'],
                           'initial_excluded':money(m['initial_excluded']),
                           'revenue_by_uf':{k:money(v) for k,v in sorted(m['by_uf'].items())},
                           'revenue_by_customer':{k:money(v) for k,v in sorted(m['by_customer'].items())},
                           'expense_by_bank':{k:money(v) for k,v in sorted(m['by_bank'].items())}})
        comparable=[m['period'] for m in result if m['comparable']]
        if not comparable: raise ValueError('Nenhum mes possui as duas fontes; base anterior preservada.')
        return {'metadata':{'generated_at':dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                            'source_file':original_name or source.name,'source_modified_at':dt.datetime.fromtimestamp(source.stat().st_mtime).isoformat(timespec='seconds'),
                            'fallback':fallback,'revenue_sheet':'RECEITA','revenue_basis':'gross',
                            'expense_basis':'competencia_registrada_ou_data_pagamento','default_period':max(comparable),
                            'last_invoice_date':last_invoice.isoformat(),'last_expense_date':last_expense.isoformat() if last_expense else None,
                            'bank_sheets':bank_sheets,'canceled_revenue_rows':canceled,'competence_date_differences':mismatch_count},
                'months':result,'records':records}
    finally: wb.close()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path)
    parser.add_argument('--original-source')
    parser.add_argument('--fallback-cache', action='store_true')
    args=parser.parse_args()
    source,fallback=(args.source,False) if args.source else choose_source()
    name=Path(args.original_source).name if args.original_source else source.name
    data=consolidate(source,name,fallback or args.fallback_cache)
    if source.resolve()!=CACHE.resolve(): shutil.copy2(source,CACHE)
    body='window.EXECUTIVO_JLE_DATA = '+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n'
    temporary=OUTPUT.with_suffix('.js.tmp');temporary.write_text(body,encoding='utf-8');temporary.replace(OUTPUT)
    print('Executivo JLE: notas',sum(m['revenue_count'] for m in data['months']),
          '| saidas',sum(m['expense_count'] for m in data['months']),'| ultimo mes comparavel',data['metadata']['default_period'])
    for m in data['months']: print(m['period'],'Receita',m['revenue'],'Despesa',m['expense'],'Resultado',m['result'])


if __name__=='__main__': main()
